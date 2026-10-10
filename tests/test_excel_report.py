"""Tests for the adaptive Excel business intelligence workbook."""

from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.insights.engine import BusinessInsight, InsightsReport
from src.reporting.excel_report import build_business_insights_excel

from src.reporting.excel_report import (
    _get_category_columns,
    _get_numeric_columns,
    _make_chart_data,
)

def sample_report():
    return InsightsReport(
        dataset_name="retail_sales.csv",
        rows_analyzed=6,
        columns_analyzed=4,
        insights=[
            BusinessInsight(
                category="Data Quality",
                severity="warning",
                title="Missing values detected",
                description="Some cells contain missing values.",
                evidence={"missing_cells": 2},
                recommendation="Review the affected columns.",
            ),
            BusinessInsight(
                category="Relationships",
                severity="info",
                title="Strong correlation detected",
                description="Two numeric columns may be related.",
                evidence={"correlation": 0.85},
                recommendation="Investigate the relationship.",
            ),
        ],
    )


def sample_data():
    return pd.DataFrame(
        {
            "Order_Date": pd.date_range("2026-01-01", periods=6, freq="D"),
            "Region": ["North", "South", "North", "West", "South", "West"],
            "Sales": [100, 200, 150, 300, 250, 350],
            "Profit": [20, 40, 25, 70, 50, 80],
        }
    )


def open_generated_workbook(excel_bytes):
    return load_workbook(BytesIO(excel_bytes), data_only=True)


def test_excel_report_returns_valid_workbook():
    result = build_business_insights_excel(
        sample_report(),
        data=sample_data(),
    )

    assert isinstance(result, bytes)
    assert result.startswith(b"PK")
    assert len(result) > 1000


def test_workbook_contains_expected_sheets():
    result = build_business_insights_excel(
        sample_report(),
        data=sample_data(),
    )

    workbook = open_generated_workbook(result)

    assert workbook.sheetnames == [
        "Executive Dashboard",
        "Statistical Summary",
        "Visual Analysis",
        "Business Insights",
        "Source Data",
    ]

    workbook.close()


def test_dashboard_contains_dataset_metrics():
    result = build_business_insights_excel(
        sample_report(),
        data=sample_data(),
    )

    workbook = open_generated_workbook(result)
    sheet = workbook["Executive Dashboard"]

    assert sheet["B5"].value == 6
    assert sheet["D5"].value == 4
    assert sheet["F5"].value == 2
    assert sheet["H5"].value == 1

    workbook.close()


def test_workbook_creates_native_excel_charts():
    result = build_business_insights_excel(
        sample_report(),
        data=sample_data(),
    )

    workbook = load_workbook(BytesIO(result), data_only=False)

    assert len(workbook["Executive Dashboard"]._charts) >= 1
    assert len(workbook["Visual Analysis"]._charts) >= 1

    workbook.close()


def test_business_insights_are_exported():
    result = build_business_insights_excel(
        sample_report(),
        data=sample_data(),
    )

    workbook = open_generated_workbook(result)
    sheet = workbook["Business Insights"]

    titles = [
        sheet.cell(row=row, column=3).value
        for row in range(5, sheet.max_row + 1)
    ]

    assert "Missing values detected" in titles
    assert "Strong correlation detected" in titles

    workbook.close()


def test_workbook_can_be_created_without_dataset():
    result = build_business_insights_excel(sample_report())

    workbook = open_generated_workbook(result)

    assert "Executive Dashboard" in workbook.sheetnames
    assert "Business Insights" in workbook.sheetnames
    assert "Source Data" in workbook.sheetnames

    workbook.close()


def test_workbook_rejects_invalid_report():
    with pytest.raises(TypeError):
        build_business_insights_excel("invalid")


def test_workbook_rejects_invalid_dataset():
    with pytest.raises(TypeError):
        build_business_insights_excel(
            sample_report(),
            data="not a DataFrame",
        )

def test_numeric_columns_prioritize_business_measures():
    data = pd.DataFrame(
        {
            "Customer ID": [101, 102, 103, 104],
            "Transaction ID": [1001, 1002, 1003, 1004],
            "Price Per Unit": [10, 20, 15, 25],
            "Total Spent": [100, 200, 150, 250],
        }
    )

    columns = _get_numeric_columns(data)

    assert columns[0] == "Total Spent"
    assert "Customer ID" not in columns
    assert "Transaction ID" not in columns


def test_category_columns_prioritize_region():
    data = pd.DataFrame(
        {
            "Customer Name": [
                "A", "B", "C", "D", "E", "F"
            ],
            "Region": [
                "North", "South", "North",
                "West", "South", "West",
            ],
            "Total Spent": [100, 200, 150, 300, 250, 350],
        }
    )

    columns = _get_category_columns(data)

    assert columns[0] == "Region"
    assert "Customer Name" in columns


def test_category_chart_uses_total_spent():
    from openpyxl import Workbook

    data = pd.DataFrame(
        {
            "Customer ID": [101, 102, 103, 104, 105, 106],
            "Region": [
                "North", "South", "North",
                "West", "South", "West",
            ],
            "Price Per Unit": [10, 20, 15, 25, 30, 35],
            "Total Spent": [100, 200, 150, 300, 250, 350],
        }
    )

    workbook = Workbook()
    worksheet = workbook.active

    sources = _make_chart_data(data, worksheet)

    assert "category" in sources

    start_row, _, columns = sources["category"]

    assert columns[0] == "Region"
    assert columns[1] == "Total Spent"
    assert worksheet.cell(start_row, 1).value == "Region"
    assert worksheet.cell(start_row, 2).value == "Total Spent"

    workbook.close()


def test_numeric_only_dataset_generates_charts_without_category():
    from openpyxl import Workbook

    data = pd.DataFrame(
        {
            "Sales": [100, 200, 150, 300, 250],
            "Profit": [20, 40, 25, 70, 50],
        }
    )

    workbook = Workbook()
    worksheet = workbook.active

    sources = _make_chart_data(data, worksheet)

    assert "distribution" in sources
    assert "scatter" in sources
    assert "category" not in sources
    assert "trend" not in sources

    workbook.close()